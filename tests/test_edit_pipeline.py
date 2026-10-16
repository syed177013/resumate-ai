from io import BytesIO

from docx import Document

from services.parser import parse_docx
from services.patcher import apply_instruction, NoChangeError, local_patch
from models import Resume, Section, Entry, Bullet


def make_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("Test Candidate")
    doc.add_paragraph("candidate@example.com | github.com/example")
    doc.add_paragraph("PROFESSIONAL SUMMARY")
    doc.add_paragraph("Python developer.")
    doc.add_paragraph("TECHNICAL SKILLS")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "SQL"
    doc.add_paragraph("SELECTED PROJECTS")
    doc.add_paragraph("Demo API")
    doc.add_paragraph("• Built a REST API")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_parser_recognizes_common_headings_and_table_content():
    resume = parse_docx(make_docx(), "test_resume.docx")
    names = [section.name.casefold() for section in resume.sections]
    assert "professional summary" in names
    assert "technical skills" in names
    assert "selected projects" in names
    skills = next(s for s in resume.sections if s.name.casefold() == "technical skills")
    assert any("Python" in bullet.text and "SQL" in bullet.text for bullet in skills.bullets)


def test_local_patch_adds_bullet():
    resume = Resume(
        id="resume1", owner="me", title="Test", sections=[
            Section(id="s1", name="Achievements", kind="bullets", bullets=[
                Bullet(id="b1", text="Existing achievement")
            ])
        ]
    )
    updated = local_patch(resume, "Add a bullet under Achievements saying improved test coverage")
    assert len(updated.sections[0].bullets) == 2
    assert updated.sections[0].bullets[-1].text == "improved test coverage"


def test_noop_local_edit_is_not_reported_as_success():
    resume = Resume(id="resume1", owner="me", title="Test", sections=[
        Section(id="s1", name="Summary", kind="text", content="Original")
    ])
    try:
        apply_instruction(resume, "Make this better", None)
    except NoChangeError:
        pass
    else:
        raise AssertionError("Expected no-op edit to raise NoChangeError")
