from __future__ import annotations

import re
from io import BytesIO
from uuid import uuid4
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from models import Resume, Section, Entry, Bullet


# Canonical names are normalized for recognition, but original heading text is kept.
KNOWN_SECTIONS = {
    "summary", "professional summary", "profile", "objective",
    "experience", "work experience", "professional experience", "employment",
    "education", "skills", "technical skills", "core skills", "projects",
    "selected projects", "certifications", "certificates", "certificates and credentials",
    "credentials", "awards", "achievements", "languages", "publications",
    "research", "research experience", "volunteer experience", "leadership",
    "leadership and activities", "interests", "coursework",
}
TEXT_SECTIONS = {"summary", "professional summary", "profile", "objective"}
ENTRY_SECTIONS = {
    "experience", "work experience", "professional experience", "employment",
    "education", "projects", "selected projects", "publications",
    "research", "research experience", "volunteer experience",
}
BULLET_SECTIONS = {
    "skills", "technical skills", "core skills", "certifications", "certificates",
    "certificates and credentials", "credentials", "awards", "achievements",
    "languages", "leadership", "leadership and activities", "interests", "coursework",
}


def uid() -> str:
    return uuid4().hex[:12]


def _normal_heading(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().rstrip(":")).casefold()


def _section_kind(name: str) -> str:
    low = _normal_heading(name)
    if low in TEXT_SECTIONS:
        return "text"
    if low in ENTRY_SECTIONS:
        return "entries"
    return "bullets"


def _iter_blocks(parent):
    """Yield paragraphs and tables in their actual document order."""
    if hasattr(parent, "element") and hasattr(parent.element, "body"):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _document_lines(doc) -> list[tuple[str, str]]:
    lines: list[tuple[str, str]] = []

    def add_paragraph(p: Paragraph):
        text = re.sub(r"\s+", " ", (p.text or "").replace("\xa0", " ")).strip()
        if text:
            lines.append((text, (p.style.name if p.style else "").casefold()))

    def walk_table(table: Table):
        for row in table.rows:
            # Keep cells in one logical line so a two-column row isn't reordered.
            cells = []
            for cell in row.cells:
                parts = []
                for paragraph in cell.paragraphs:
                    value = re.sub(r"\s+", " ", (paragraph.text or "").replace("\xa0", " ")).strip()
                    if value:
                        parts.append(value)
                cell_text = " ".join(parts).strip()
                if cell_text and (not cells or cell_text != cells[-1]):
                    cells.append(cell_text)
            if cells:
                lines.append((" | ".join(cells), "table"))
            # Include nested tables rather than silently discarding them.
            for cell in row.cells:
                for block in _iter_blocks(cell):
                    if isinstance(block, Table):
                        walk_table(block)

    for block in _iter_blocks(doc):
        if isinstance(block, Paragraph):
            add_paragraph(block)
        else:
            walk_table(block)
    return lines


def _is_heading(text: str, style: str) -> bool:
    normalized = _normal_heading(text)
    if "heading" in style or normalized in KNOWN_SECTIONS:
        return True
    # Some resume templates use bold/all-caps headings with Normal style.
    stripped = text.strip()
    letters = [c for c in stripped if c.isalpha()]
    if letters and len(stripped) <= 48 and len(stripped.split()) <= 6:
        if stripped == stripped.upper() and normalized not in {
            "python", "java", "sql", "c++", "r", "machine learning",
        }:
            return True
    return False


def _is_bullet(text: str, style: str) -> bool:
    return (
        text.startswith(("•", "●", "▪", "◦", "-", "*", "–", "—"))
        or "list" in style
    )


def _clean_bullet(text: str) -> str:
    return re.sub(r"^[\s•●▪◦*\-–—]+", "", text).strip()


def parse_docx(content: bytes, filename: str, owner: str = "me") -> Resume:
    doc = Document(BytesIO(content))
    lines = _document_lines(doc)
    title = filename.rsplit(".", 1)[0].replace("_", " ").title()
    name = lines[0][0] if lines else None
    contact: dict[str, str] = {}

    # Keep the common one-line contact row. Don't throw away other header lines.
    header_skip: set[int] = {0} if lines else set()
    if len(lines) > 1 and ("@" in lines[1][0] or "|" in lines[1][0] or re.search(r"\bhttps?://|\blinkedin\.com\b|\bgithub\.com\b", lines[1][0], re.I)):
        contact["details"] = lines[1][0]
        header_skip.add(1)

    sections: list[Section] = []
    current: Section | None = None
    current_entry: Entry | None = None

    for i, (text, style) in enumerate(lines):
        if i in header_skip:
            continue

        if _is_heading(text, style):
            label = text.rstrip(":").strip()
            kind = _section_kind(label)
            current = Section(
                id=uid(), name=label, kind=kind,
                content="" if kind == "text" else None,
                entries=[] if kind == "entries" else None,
                bullets=[] if kind == "bullets" else None,
            )
            sections.append(current)
            current_entry = None
            continue

        if current is None:
            current = Section(id=uid(), name="Summary", kind="text", content="")
            sections.append(current)

        if current.kind == "text":
            current.content = " ".join(filter(None, [current.content, text])).strip()
        elif current.kind == "bullets":
            current.bullets.append(Bullet(id=uid(), text=_clean_bullet(text)))
        elif _is_bullet(text, style):
            if current_entry is None:
                current_entry = Entry(id=uid(), title="Details", bullets=[])
                current.entries.append(current_entry)
            current_entry.bullets.append(Bullet(id=uid(), text=_clean_bullet(text)))
        else:
            parts = [part.strip() for part in text.split("|", 1)]
            current_entry = Entry(
                id=uid(),
                title=parts[0],
                org=parts[1] if len(parts) > 1 else None,
                bullets=[],
            )
            current.entries.append(current_entry)

    if not sections:
        sections = [Section(id=uid(), name="Summary", kind="text", content="")]
    return Resume(
        id=uuid4().hex,
        owner=owner if owner in {"me", "friend"} else "me",
        title=title,
        name=name,
        contact=contact,
        sections=sections,
    )
